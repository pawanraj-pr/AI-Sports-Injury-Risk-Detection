import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="14181B", end_color="14181B", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _style_header_row(ws, row_idx, num_cols):
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="left")


def _autofit(ws):
    for col_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(length + 2, 10), 45)


def build_summary_report_excel(athlete, reports_with_videos) -> bytes:
    wb = Workbook()

    # ---------- Summary sheet ----------
    ws = wb.active
    ws.title = "Summary"

    ws["A1"] = "Combined Biomechanics Report"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Athlete: {athlete.athlete_code}  |  Sport: {athlete.sport_type}"
    ws["A3"] = f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"
    ws["A4"] = f"Videos analyzed: {len(reports_with_videos)}"

    def avg_field(field):
        vals = [getattr(r, field) for _, r in reports_with_videos if getattr(r, field) is not None]
        return round(sum(vals) / len(vals), 1) if vals else None

    metrics = [
        ("Average Movement Quality Score", avg_field("movement_quality_score")),
        ("Average Hip Stability", avg_field("hip_stability_score")),
        ("Average Landing Mechanics", avg_field("landing_mechanics_score")),
        ("Average Joint Alignment", avg_field("joint_alignment_score")),
        ("Average Balance", avg_field("balance_score")),
        ("Average Movement Symmetry", avg_field("movement_symmetry_score")),
        ("Average Knee Valgus Score", avg_field("knee_valgus_score")),
    ]

    row = 6
    ws.cell(row=row, column=1, value="Metric")
    ws.cell(row=row, column=2, value="Value")
    _style_header_row(ws, row, 2)
    for label, value in metrics:
        row += 1
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=value if value is not None else "-")

    # ---------- Per-video sheet ----------
    ws2 = wb.create_sheet("Videos")
    headers = ["Video", "Activity", "Uploaded", "Movement Quality", "Risk Category",
               "Left Knee Angle", "Right Knee Angle", "Knee Valgus", "Hip Stability",
               "Landing Mechanics", "Joint Alignment", "Balance", "Symmetry", "Trunk Lean"]
    for col, h in enumerate(headers, start=1):
        ws2.cell(row=1, column=col, value=h)
    _style_header_row(ws2, 1, len(headers))

    for i, (video, report) in enumerate(
        sorted(reports_with_videos, key=lambda vr: vr[0].uploaded_at), start=2
    ):
        values = [
            video.filename, video.activity_type.value.replace("_", " "),
            video.uploaded_at.strftime("%Y-%m-%d"),
            report.movement_quality_score, report.risk_category,
            report.avg_left_knee_angle, report.avg_right_knee_angle,
            report.knee_valgus_score, report.hip_stability_score,
            report.landing_mechanics_score, report.joint_alignment_score,
            report.balance_score, report.movement_symmetry_score,
            report.trunk_lean_degrees,
        ]
        for col, val in enumerate(values, start=1):
            ws2.cell(row=i, column=col, value=val if val is not None else "-")

    _autofit(ws)
    _autofit(ws2)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
